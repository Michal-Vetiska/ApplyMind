"""Knowledge document ingestion: extract text, chunk, embed, persist."""

from __future__ import annotations

import hashlib
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.logging import get_logger
from app.integrations.llm import get_llm_client
from app.models.knowledge import KnowledgeChunk, KnowledgeDocument
from app.modules.rag.chunking import chunk_text
from app.modules.rag.embeddings import embed_texts

logger = get_logger(__name__)


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text)
    raise ValueError(f"Unsupported knowledge file type: {suffix}")


def file_checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


async def ingest_document(
    db: AsyncSession,
    *,
    title: str,
    document_type: str,
    file_path: str | None = None,
    text: str | None = None,
    source: str = "upload",
    scope: str = "PERSONAL_KNOWLEDGE",
    approved: bool = False,
    metadata: dict | None = None,
    reindex: bool = False,
) -> KnowledgeDocument:
    settings = get_settings()
    content = text or ""
    checksum = hashlib.sha256(content.encode("utf-8")).hexdigest() if content else ""

    if file_path:
        path = Path(file_path)
        if not path.is_absolute():
            path = Path(settings.data_dir) / path
        content = extract_text_from_file(path)
        checksum = file_checksum(path)

    if not content.strip():
        raise ValueError("Document has no extractable text")

    existing = None
    if checksum:
        result = await db.execute(
            select(KnowledgeDocument).where(
                KnowledgeDocument.checksum == checksum,
                KnowledgeDocument.deleted_at.is_(None),
            )
        )
        existing = result.scalar_one_or_none()

    if existing and not reindex:
        return existing

    doc = existing
    if doc is None:
        doc = KnowledgeDocument(
            title=title,
            document_type=document_type,
            source=source,
            scope=scope,
            file_path=file_path,
            checksum=checksum,
            approved=approved,
            metadata_json=metadata or {},
        )
        db.add(doc)
        await db.flush()
    else:
        doc.title = title
        doc.document_type = document_type
        doc.scope = scope
        doc.approved = approved
        doc.metadata_json = metadata or doc.metadata_json
        await db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == doc.id))

    chunks = chunk_text(content)
    llm = get_llm_client(db)
    embeddings: list[list[float]] = []
    if llm.is_configured:
        try:
            embeddings = await embed_texts([c.content for c in chunks], client=llm)
        except Exception as exc:
            logger.warning("embedding_failed", error=str(exc))
            embeddings = []
    else:
        logger.info("embeddings_skipped_llm_not_configured")

    for i, chunk in enumerate(chunks):
        emb = embeddings[i] if i < len(embeddings) else None
        db.add(
            KnowledgeChunk(
                document_id=doc.id,
                content=chunk.content,
                embedding=emb,
                metadata_json=chunk.metadata or {},
                chunk_index=chunk.chunk_index,
                token_count=chunk.token_count,
                section_heading=chunk.section_heading,
            )
        )

    doc.indexed_at = datetime.now(UTC)
    doc.embedding_model = settings.llm_embedding_model if embeddings else None
    await db.flush()
    return doc


async def reindex_document(db: AsyncSession, document_id: str) -> KnowledgeDocument:
    doc = await db.get(KnowledgeDocument, document_id)
    if doc is None or doc.deleted_at is not None:
        raise ValueError("Document not found")
    text = None
    if doc.file_path:
        return await ingest_document(
            db,
            title=doc.title,
            document_type=doc.document_type,
            file_path=doc.file_path,
            source=doc.source,
            scope=doc.scope,
            approved=doc.approved,
            metadata=doc.metadata_json,
            reindex=True,
        )
    # Fall back to concatenating existing chunk text before delete
    result = await db.execute(
        select(KnowledgeChunk)
        .where(KnowledgeChunk.document_id == doc.id)
        .order_by(KnowledgeChunk.chunk_index)
    )
    chunks = list(result.scalars().all())
    text = "\n\n".join(c.content for c in chunks)
    return await ingest_document(
        db,
        title=doc.title,
        document_type=doc.document_type,
        text=text,
        source=doc.source,
        scope=doc.scope,
        approved=doc.approved,
        metadata=doc.metadata_json,
        reindex=True,
    )
